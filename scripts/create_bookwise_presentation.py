from __future__ import annotations

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_THEME_COLOR
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "BookWise_Project_Presentation.pptx"

SLIDE_W = Inches(20)
SLIDE_H = Inches(11.25)

COLORS = {
    "ink": "211A13",
    "ink2": "3D342A",
    "ivory": "F8F4EC",
    "paper": "FFFCF7",
    "line": "E8DFC8",
    "muted": "7A6D5A",
    "gold": "C5A059",
    "gold2": "D4AF37",
    "burgundy": "722F37",
    "green": "2D6A4F",
    "blue": "3B6178",
    "sand": "FAF5EB",
    "lavender": "EEE9F6",
}


def rgb(name: str) -> RGBColor:
    return RGBColor.from_string(COLORS[name] if name in COLORS else name)


def set_fill(shape, color: str, transparency: int = 0):
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    shape.fill.transparency = transparency


def set_line(shape, color: str = "line", width: float = 1.0, transparency: int = 0):
    shape.line.color.rgb = rgb(color)
    shape.line.width = Pt(width)
    shape.line.transparency = transparency


def add_text(slide, text, x, y, w, h, *, size=18, color="ink", bold=False,
             font="Aptos", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
             italic=False, margin=0.05, fit=True):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = str(text)
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    if fit:
        try:
            tf.fit_text(font_family=font, max_size=Pt(size))
        except Exception:
            pass
    return box


def add_rich_text(slide, runs, x, y, w, h, *, size=18, color="ink", font="Aptos",
                  align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP, margin=0.08):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    for item in runs:
        run = p.add_run()
        run.text = item.get("text", "")
        run.font.name = item.get("font", font)
        run.font.size = Pt(item.get("size", size))
        run.font.bold = item.get("bold", False)
        run.font.italic = item.get("italic", False)
        run.font.color.rgb = rgb(item.get("color", color))
    return box


def add_box(slide, x, y, w, h, *, fill="paper", line="line", radius=True,
            transparency=0):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    set_fill(shape, fill, transparency)
    set_line(shape, line)
    return shape


def add_line(slide, x1, y1, x2, y2, *, color="gold", width=2.0, arrow=False):
    line = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    set_line(line, color, width)
    if arrow:
        line.line.end_arrowhead = True
    return line


def add_pill(slide, text, x, y, w, *, fill="sand", color="muted", size=10, line="line"):
    add_box(slide, x, y, w, 0.36, fill=fill, line=line, radius=True)
    add_text(slide, text, x + 0.08, y + 0.03, w - 0.16, 0.28, size=size, color=color,
             bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE, margin=0)


def add_bullets(slide, items, x, y, w, h, *, size=18, color="ink", bullet_color="gold",
                gap=0.12, font="Aptos"):
    cursor = y
    for item in items:
        if isinstance(item, tuple):
            label, body = item
            add_text(slide, "•", x, cursor, 0.25, 0.3, size=size + 2, color=bullet_color, bold=True, margin=0)
            add_rich_text(slide, [
                {"text": label, "bold": True, "color": color},
                {"text": body, "color": color},
            ], x + 0.28, cursor, w - 0.28, 0.55, size=size, color=color, margin=0)
        else:
            add_text(slide, "•", x, cursor, 0.25, 0.3, size=size + 2, color=bullet_color, bold=True, margin=0)
            add_text(slide, item, x + 0.28, cursor, w - 0.28, 0.55, size=size, color=color, margin=0)
        cursor += gap + 0.52


def add_metric(slide, value, label, x, y, w, h, *, accent="gold", note=None):
    add_box(slide, x, y, w, h, fill="paper", line="line")
    add_text(slide, value, x + 0.18, y + 0.16, w - 0.36, 0.55, size=28, color=accent, bold=True,
             font="Aptos Display", valign=MSO_ANCHOR.MIDDLE, margin=0)
    add_text(slide, label, x + 0.18, y + 0.78, w - 0.36, 0.38, size=12, color="ink", bold=True, margin=0)
    if note:
        add_text(slide, note, x + 0.18, y + 1.14, w - 0.36, h - 1.22, size=9, color="muted", margin=0)


def add_title(slide, number, eyebrow, title, subtitle=None):
    add_text(slide, str(number), 0.65, 0.47, 0.55, 0.38, size=13, color="gold", bold=True, margin=0)
    add_text(slide, eyebrow.upper(), 1.28, 0.47, 7.2, 0.28, size=10, color="muted", bold=True, margin=0)
    add_text(slide, title, 0.65, 0.92, 18.7, 0.68, size=28, color="ink", bold=True,
             font="Aptos Display", margin=0)
    if subtitle:
        add_text(slide, subtitle, 0.68, 1.63, 18.0, 0.4, size=12, color="muted", margin=0)
    add_line(slide, 0.65, 2.12, 19.35, 2.12, color="line", width=1)


def add_footer(slide, text="BookWise • Intelligent Book Discovery"):
    add_line(slide, 0.65, 10.82, 19.35, 10.82, color="line", width=0.8)
    add_text(slide, text, 0.68, 10.9, 10.0, 0.2, size=8, color="muted", margin=0)


def add_slide(prs, *, bg="ivory", number=None, eyebrow=None, title=None, subtitle=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    background = slide.background.fill
    background.solid()
    background.fore_color.rgb = rgb(bg)
    if number is not None:
        add_title(slide, number, eyebrow or "BookWise", title or "", subtitle)
        add_footer(slide)
    return slide


def add_card_header(slide, label, x, y, w, *, color="gold", icon=None):
    if icon:
        add_text(slide, icon, x, y, 0.32, 0.32, size=16, color=color, bold=True, margin=0)
        x += 0.4
        w -= 0.4
    add_text(slide, label.upper(), x, y + 0.03, w, 0.25, size=9, color=color, bold=True, margin=0)


def add_book_card(slide, x, y, w, h, title, author, score, accent="gold"):
    add_box(slide, x, y, w, h, fill="paper", line="line")
    cover = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + 0.18), Inches(y + 0.18), Inches(w * 0.31), Inches(h - 0.36))
    set_fill(cover, accent)
    set_line(cover, accent)
    add_text(slide, "BOOK\nWISE", x + 0.25, y + 0.47, w * 0.17, 0.52, size=11, color="paper", bold=True,
             font="Aptos Display", align=PP_ALIGN.CENTER, margin=0)
    add_text(slide, title, x + w * 0.37, y + 0.26, w * 0.55, 0.75, size=14, color="ink", bold=True,
             font="Aptos Display", margin=0)
    add_text(slide, author, x + w * 0.37, y + 1.08, w * 0.55, 0.32, size=10, color="muted", margin=0)
    add_pill(slide, score, x + w * 0.37, y + h - 0.58, w * 0.43, fill="sand", color=accent, size=9, line="line")


def add_process_step(slide, x, y, n, title, body, *, color="gold", w=2.0, h=1.25):
    add_box(slide, x, y, w, h, fill="paper", line="line")
    add_text(slide, f"{n:02d}", x + 0.16, y + 0.14, 0.42, 0.28, size=11, color=color, bold=True, margin=0)
    add_text(slide, title, x + 0.16, y + 0.46, w - 0.32, 0.32, size=13, color="ink", bold=True, margin=0)
    add_text(slide, body, x + 0.16, y + 0.82, w - 0.32, h - 0.94, size=9, color="muted", margin=0)


def add_bar(slide, label, value, max_value, x, y, w, *, color="gold", suffix=""):
    add_text(slide, label, x, y, w * 0.43, 0.28, size=10, color="ink", margin=0)
    track_x = x + w * 0.44
    track_w = w * 0.44
    add_box(slide, track_x, y + 0.04, track_w, 0.18, fill="line", line="line", radius=True)
    fill_w = max(0.05, track_w * (value / max_value if max_value else 0))
    add_box(slide, track_x, y + 0.04, fill_w, 0.18, fill=color, line=color, radius=True)
    add_text(slide, f"{value:.4f}{suffix}", track_x + track_w + 0.1, y - 0.02, w * 0.12, 0.28, size=10, color=color, bold=True, margin=0)


def add_table(slide, headers, rows, x, y, widths, row_h=0.45, *, header_fill="ink", font_size=9):
    current_x = x
    for header, width in zip(headers, widths):
        cell = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(current_x), Inches(y), Inches(width), Inches(row_h))
        set_fill(cell, header_fill)
        set_line(cell, header_fill)
        add_text(slide, header, current_x + 0.06, y + 0.09, width - 0.12, row_h - 0.1, size=font_size, color="paper", bold=True, margin=0)
        current_x += width
    for r, row in enumerate(rows):
        current_x = x
        fill = "paper" if r % 2 == 0 else "sand"
        for c, (value, width) in enumerate(zip(row, widths)):
            cell = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(current_x), Inches(y + row_h * (r + 1)), Inches(width), Inches(row_h))
            set_fill(cell, fill)
            set_line(cell, "line", 0.6)
            color = "green" if r == len(rows) - 1 and c == 0 else "ink"
            add_text(slide, value, current_x + 0.06, y + row_h * (r + 1) + 0.09, width - 0.12, row_h - 0.1,
                     size=font_size, color=color, bold=(r == len(rows) - 1 and c == 0), margin=0)
            current_x += width


def slide_01(prs):
    slide = add_slide(prs, bg="ink")
    add_text(slide, "RECOMMENDER SYSTEMS", 0.8, 0.72, 8.0, 0.3, size=12, color="gold2", bold=True, margin=0)
    add_text(slide, "BookWise", 0.8, 1.48, 12.0, 1.0, size=46, color="paper", bold=True, font="Aptos Display", margin=0)
    add_text(slide, "Intelligent Book Discovery & Recommendation", 0.85, 2.62, 14.5, 0.52, size=24, color="gold2", bold=True, font="Aptos Display", margin=0)
    add_text(slide, "A reader searches a 10,000-book catalogue. BookWise combines popularity, content similarity, reader patterns and personal feedback to surface the next great read.", 0.85, 3.42, 10.8, 1.0, size=16, color="ivory", margin=0)
    add_pill(slide, "React + Flask + SQLAlchemy", 0.85, 4.72, 2.55, fill="ink2", color="gold2", line="gold")
    add_pill(slide, "4 recommendation engines", 3.58, 4.72, 2.25, fill="ink2", color="gold2", line="gold")
    add_pill(slide, "Goodbooks-10k data", 6.02, 4.72, 1.95, fill="ink2", color="gold2", line="gold")

    # Editorial-style mock of the running app.
    add_box(slide, 12.1, 1.2, 6.8, 7.2, fill="paper", line="gold")
    add_text(slide, "BOOKWISE", 12.48, 1.56, 2.0, 0.25, size=10, color="muted", bold=True, margin=0)
    add_text(slide, "Your next chapter", 12.48, 2.05, 4.9, 0.45, size=24, color="ink", bold=True, font="Aptos Display", margin=0)
    add_text(slide, "Curated shelves shaped by your taste.", 12.48, 2.58, 4.5, 0.28, size=11, color="muted", margin=0)
    add_line(slide, 12.48, 3.1, 18.45, 3.1, color="line", width=1)
    add_text(slide, "RECOMMENDED FOR YOU", 12.48, 3.42, 3.0, 0.22, size=9, color="gold", bold=True, margin=0)
    add_book_card(slide, 12.48, 3.82, 2.55, 2.75, "The Night Circus", "Erin Morgenstern", "94% match", "burgundy")
    add_book_card(slide, 15.25, 3.82, 2.55, 2.75, "Dune", "Frank Herbert", "91% match", "blue")
    add_pill(slide, "Why this book", 12.48, 6.95, 1.28, fill="sand", color="gold", size=8)
    add_text(slide, "Shared themes • Similar readers • Fresh discovery", 13.9, 7.01, 4.3, 0.2, size=9, color="muted", margin=0)
    add_text(slide, "Project presentation", 0.85, 9.55, 4.0, 0.28, size=12, color="gold2", bold=True, margin=0)
    add_text(slide, "BookWise web app • Academic recommender systems project", 0.85, 9.9, 8.4, 0.28, size=10, color="muted", margin=0)
    add_footer(slide, "BookWise • Intelligent Book Discovery")


def slide_02(prs):
    slide = add_slide(prs, number=2, eyebrow="Problem and Motivation", title="Finding a good book is easy. Finding the right next book is not.", subtitle="BookWise is designed around the decisions readers actually make.")
    add_box(slide, 0.8, 2.65, 8.4, 6.9, fill="ink", line="ink")
    add_text(slide, "A catalogue can be large without being helpful.", 1.2, 3.08, 7.3, 0.75, size=24, color="paper", bold=True, font="Aptos Display", margin=0)
    add_bullets(slide, [
        ("Choice overload. ", "Thousands of titles make browsing feel like work."),
        ("One-size-fits-all lists. ", "Global popularity does not explain personal fit."),
        ("Sparse signals. ", "New readers have little history for a collaborative model."),
        ("Low trust. ", "A recommendation without a reason is hard to act on."),
    ], 1.2, 4.15, 7.2, 3.2, size=16, color="ivory", bullet_color="gold2", gap=0.2)
    add_pill(slide, "Reader problem", 1.2, 8.55, 1.45, fill="ink2", color="gold2", line="gold")
    add_text(slide, "I want to discover something I will actually finish.", 2.82, 8.61, 5.7, 0.25, size=13, color="paper", italic=True, margin=0)

    add_text(slide, "What the reader needs", 10.0, 2.82, 4.4, 0.35, size=14, color="muted", bold=True, margin=0)
    needs = [
        ("Relevant", "Books aligned with themes, genres and authors I already enjoy."),
        ("Serendipitous", "New titles that similar readers discovered."),
        ("Explainable", "A short, human-readable reason for every shelf."),
        ("Personal", "A place to rate, save, track and share reading choices."),
    ]
    for i, (label, body) in enumerate(needs):
        y = 3.45 + i * 1.35
        add_box(slide, 10.0, y, 8.7, 1.05, fill="paper", line="line")
        add_text(slide, f"0{i+1}", 10.25, y + 0.22, 0.45, 0.3, size=13, color="gold", bold=True, margin=0)
        add_text(slide, label, 10.9, y + 0.18, 2.3, 0.32, size=15, color="ink", bold=True, margin=0)
        add_text(slide, body, 13.2, y + 0.18, 5.1, 0.45, size=11, color="muted", margin=0)


def slide_03(prs):
    slide = add_slide(prs, number=3, eyebrow="Objectives and Solution", title="From catalogue search to a personal reading system", subtitle="The product combines discovery, personalization and reader-owned tools in one flow.")
    add_box(slide, 0.8, 2.62, 8.4, 7.35, fill="paper", line="line")
    add_card_header(slide, "Objectives", 1.2, 3.0, 3.0, color="gold")
    objectives = [
        "Search and browse a structured 10,000-book catalogue",
        "Rank recommendations with multiple complementary models",
        "Make cold-start behavior useful from the first visit",
        "Explain why a title was recommended",
        "Support private reading workflows and shareable lists",
    ]
    add_bullets(slide, objectives, 1.2, 3.55, 7.2, 4.0, size=15, color="ink", bullet_color="gold", gap=0.14)
    add_pill(slide, "Outcome", 1.2, 8.55, 1.0, fill="sand", color="gold", size=9)
    add_text(slide, "A discovery experience that improves as the reader interacts.", 2.35, 8.59, 5.9, 0.25, size=12, color="muted", italic=True, margin=0)

    add_text(slide, "What BookWise returns", 10.0, 2.9, 5.0, 0.35, size=14, color="muted", bold=True, margin=0)
    outputs = [
        ("Browse", "Popular shelves, hidden gems, genre filters and search."),
        ("Personalize", "For You modes, preferences, feedback and taste clusters."),
        ("Explain", "Theme, author, similarity and similar-reader signals."),
        ("Read", "Ratings, shelves, progress, notes, reviews and lists."),
        ("Inspect", "Offline metrics, genre prediction and ML insights."),
    ]
    for i, (label, body) in enumerate(outputs):
        y = 3.45 + i * 1.15
        add_box(slide, 10.0, y, 8.75, 0.92, fill="paper", line="line")
        add_text(slide, label, 10.28, y + 0.19, 1.55, 0.28, size=14, color="ink", bold=True, margin=0)
        add_text(slide, body, 11.9, y + 0.19, 6.4, 0.3, size=11, color="muted", margin=0)
    add_line(slide, 9.45, 4.0, 9.85, 4.0, color="gold", width=2, arrow=True)
    add_text(slide, "reader signal", 8.9, 3.62, 1.1, 0.2, size=8, color="gold", align=PP_ALIGN.CENTER, margin=0)


def slide_04(prs):
    slide = add_slide(prs, number=4, eyebrow="Recommender Fundamentals", title="Four models, one reader experience", subtitle="No single algorithm is best at every stage of the reader journey.")
    add_box(slide, 0.8, 2.65, 18.4, 1.15, fill="ink", line="ink")
    add_text(slide, "User profile", 1.25, 2.98, 1.4, 0.28, size=12, color="gold2", bold=True, margin=0)
    add_text(slide, "ratings • favorites • genres • feedback • reading progress", 2.85, 2.98, 6.0, 0.28, size=13, color="paper", margin=0)
    add_text(slide, "→", 9.25, 2.9, 0.55, 0.38, size=22, color="gold2", bold=True, align=PP_ALIGN.CENTER, margin=0)
    add_text(slide, "ranked candidates", 10.05, 2.98, 1.65, 0.28, size=12, color="gold2", bold=True, margin=0)
    add_text(slide, "books • scores • explanations • diversity", 11.95, 2.98, 5.8, 0.28, size=13, color="paper", margin=0)

    models = [
        ("1", "Popularity", "Weighted rating baseline", "Reliable from zero signals", "gold"),
        ("2", "Content-based", "TF-IDF + cosine similarity", "Themes, authors and explainability", "burgundy"),
        ("3", "Collaborative", "SVD latent factors", "Serendipity from similar readers", "blue"),
        ("4", "Hybrid", "Gated ensemble", "Balances fit, novelty and cold start", "green"),
    ]
    for i, (num, label, method, best_for, accent) in enumerate(models):
        x = 0.8 + i * 4.65
        add_box(slide, x, 4.45, 4.2, 3.85, fill="paper", line="line")
        add_text(slide, num, x + 0.3, 4.78, 0.5, 0.55, size=24, color=accent, bold=True, font="Aptos Display", margin=0)
        add_text(slide, label, x + 0.95, 4.8, 2.8, 0.38, size=17, color="ink", bold=True, font="Aptos Display", margin=0)
        add_line(slide, x + 0.3, 5.55, x + 3.85, 5.55, color="line", width=1)
        add_text(slide, method, x + 0.3, 5.85, 3.55, 0.3, size=12, color="ink", bold=True, margin=0)
        add_text(slide, best_for, x + 0.3, 6.45, 3.45, 0.62, size=12, color="muted", margin=0)
        add_pill(slide, "model output", x + 0.3, 7.45, 1.25, fill="sand", color=accent, size=8)
        add_text(slide, "rank + reason", x + 1.72, 7.51, 1.75, 0.2, size=10, color="muted", margin=0)

    add_text(slide, "Content-based filtering gives a new reader a strong starting point; collaborative filtering becomes more useful as ratings accumulate.", 1.0, 9.05, 17.8, 0.42, size=14, color="ink", italic=True, align=PP_ALIGN.CENTER, margin=0)


def slide_05(prs):
    slide = add_slide(prs, number=5, eyebrow="System Architecture", title="A full-stack pipeline from data to discovery", subtitle="The same architecture supports public browsing, personalized feeds, and reader-owned tools.")
    layers = [
        ("React / Vite UI", "Home • For You • Library • ML Insights", "gold"),
        ("Flask API", "books • recommendations • discovery • reader", "burgundy"),
        ("Recommendation layer", "Popularity • TF-IDF • SVD • Hybrid", "blue"),
        ("Data + artifacts", "Goodbooks-10k • ratings • models • metrics", "green"),
    ]
    for i, (label, body, accent) in enumerate(layers):
        y = 2.72 + i * 1.48
        add_box(slide, 1.2, y, 6.8, 1.05, fill="paper", line="line")
        add_text(slide, label, 1.55, y + 0.2, 2.7, 0.3, size=15, color=accent, bold=True, font="Aptos Display", margin=0)
        add_text(slide, body, 4.25, y + 0.2, 3.35, 0.3, size=11, color="muted", margin=0)
        if i < len(layers) - 1:
            add_line(slide, 4.6, y + 1.05, 4.6, y + 1.42, color="gold", width=1.5, arrow=True)

    add_box(slide, 9.0, 2.72, 9.65, 6.98, fill="ink", line="gold")
    add_text(slide, "REQUEST → RANK → EXPLAIN", 9.45, 3.1, 5.2, 0.3, size=12, color="gold2", bold=True, margin=0)
    flow = [
        ("1", "Reader action", "Search, rate, save, follow, or open a shelf"),
        ("2", "Candidate generation", "Retrieve popular, similar, or reader-pattern candidates"),
        ("3", "Scoring", "Blend model scores, novelty and diversity constraints"),
        ("4", "Explanation", "Return reasons, mode, algorithm and confidence signals"),
        ("5", "Feedback loop", "Clicks, likes and dismissals refine future shelves"),
    ]
    for i, (num, label, body) in enumerate(flow):
        y = 3.82 + i * 1.0
        add_text(slide, num, 9.5, y, 0.35, 0.25, size=11, color="gold2", bold=True, margin=0)
        add_text(slide, label, 10.05, y, 2.3, 0.25, size=13, color="paper", bold=True, margin=0)
        add_text(slide, body, 12.55, y, 5.5, 0.3, size=10, color="ivory", margin=0)
        if i < len(flow) - 1:
            add_line(slide, 9.67, y + 0.36, 9.67, y + 0.83, color="gold", width=1)


def slide_06(prs):
    slide = add_slide(prs, number=6, eyebrow="Data and Ingestion", title="The model sees a large, sparse reading graph", subtitle="BookWise uses the Goodbooks-10k catalogue and an 80/20 train–test split for offline evaluation.")
    add_metric(slide, "10,000", "books in the catalogue", 0.9, 2.7, 3.35, 1.62, accent="burgundy")
    add_metric(slide, "53,424", "readers represented", 4.55, 2.7, 3.35, 1.62, accent="blue")
    add_metric(slide, "5.98M", "ratings in the dataset", 8.2, 2.7, 3.35, 1.62, accent="gold")
    add_metric(slide, "98.88%", "rating-matrix sparsity", 11.85, 2.7, 3.35, 1.62, accent="green")
    add_metric(slide, "1.12%", "matrix density", 15.5, 2.7, 3.35, 1.62, accent="burgundy")

    add_box(slide, 0.9, 4.8, 8.6, 4.85, fill="paper", line="line")
    add_card_header(slide, "Data flow", 1.3, 5.18, 2.0, color="gold")
    steps = [
        ("CSV catalogue", "books, authors, tags, genres"),
        ("Train / test split", "4.78M train • 1.20M held out"),
        ("Model artifacts", "TF-IDF matrix, vectorizer, SVD factors"),
        ("API responses", "ranked books + explanations + metrics"),
    ]
    for i, (label, body) in enumerate(steps):
        y = 5.75 + i * 0.82
        add_text(slide, f"0{i+1}", 1.35, y, 0.4, 0.25, size=11, color="gold", bold=True, margin=0)
        add_text(slide, label, 1.95, y, 2.3, 0.25, size=12, color="ink", bold=True, margin=0)
        add_text(slide, body, 4.35, y, 4.4, 0.25, size=11, color="muted", margin=0)
        if i < len(steps) - 1:
            add_line(slide, 1.53, y + 0.33, 1.53, y + 0.68, color="gold", width=1, arrow=True)

    add_box(slide, 9.9, 4.8, 8.95, 4.85, fill="ink", line="gold")
    add_card_header(slide, "Why sparsity matters", 10.35, 5.18, 2.5, color="gold2")
    add_text(slide, "98.88%", 10.35, 5.9, 3.0, 0.8, size=34, color="gold2", bold=True, font="Aptos Display", margin=0)
    add_text(slide, "of the user × book matrix is empty", 10.38, 6.65, 4.2, 0.28, size=13, color="paper", margin=0)
    add_bullets(slide, [
        ("Pure CF is brittle. ", "A new or lightly active reader has little overlap with the graph."),
        ("Metadata adds a bridge. ", "Themes, authors and genres provide inductive bias."),
        ("Hybrid is safer. ", "The system falls back to popularity and content when needed."),
    ], 10.35, 7.25, 7.75, 1.9, size=12, color="ivory", bullet_color="gold2", gap=0.1)


def slide_07(prs):
    slide = add_slide(prs, number=7, eyebrow="Content-Based Model", title="Books are represented by the language around them", subtitle="TF-IDF turns metadata into vectors; cosine similarity finds books with overlapping signals.")
    add_box(slide, 0.85, 2.65, 8.5, 6.95, fill="paper", line="line")
    add_card_header(slide, "Metadata soup", 1.25, 3.05, 2.3, color="burgundy")
    fields = [
        ("Title", "The Left Hand of Darkness"),
        ("Author", "Ursula K. Le Guin"),
        ("Tags", "science fiction • space • politics • identity"),
        ("Genre", "Science Fiction"),
    ]
    for i, (label, value) in enumerate(fields):
        y = 3.65 + i * 0.78
        add_text(slide, label, 1.3, y, 1.15, 0.24, size=10, color="muted", bold=True, margin=0)
        add_box(slide, 2.55, y - 0.04, 5.95, 0.45, fill="sand", line="line")
        add_text(slide, value, 2.75, y + 0.08, 5.55, 0.2, size=11, color="ink", margin=0)
    add_line(slide, 1.3, 7.05, 8.5, 7.05, color="line", width=1)
    add_text(slide, "vector", 1.3, 7.35, 0.8, 0.25, size=10, color="burgundy", bold=True, margin=0)
    add_text(slide, "[0.00, 0.18, 0.42, 0.00, 0.77, …]", 2.55, 7.35, 4.5, 0.25, size=12, color="ink", font="Courier New", margin=0)
    add_text(slide, "A user profile is the rating-weighted centroid of books they enjoyed.", 1.3, 8.2, 7.1, 0.5, size=13, color="muted", italic=True, margin=0)

    add_box(slide, 9.9, 2.65, 8.95, 6.95, fill="ink", line="gold")
    add_card_header(slide, "Why this is explainable", 10.35, 3.05, 3.3, color="gold2")
    add_text(slide, "similarity = shared evidence", 10.35, 3.55, 6.2, 0.38, size=22, color="paper", bold=True, font="Aptos Display", margin=0)
    add_text(slide, "The API can point to the signals that connected two books:", 10.35, 4.15, 6.6, 0.28, size=12, color="ivory", margin=0)
    for i, text in enumerate(["same author", "same primary genre", "common reader tags", "high lexical similarity"]):
        y = 4.78 + i * 0.65
        add_pill(slide, text, 10.35, y, 2.55 if i != 3 else 2.85, fill="ink2", color="gold2", size=9, line="gold")
    add_text(slide, "Real artifact example", 14.05, 4.84, 2.3, 0.22, size=10, color="gold2", bold=True, margin=0)
    add_text(slide, "cosine 0.9167", 14.05, 5.15, 2.6, 0.45, size=28, color="gold2", bold=True, font="Aptos Display", margin=0)
    add_text(slide, "top precomputed similarity score in the project lookup table", 14.05, 5.7, 3.6, 0.55, size=11, color="ivory", margin=0)
    add_line(slide, 10.35, 7.5, 18.2, 7.5, color="ink2", width=1)
    add_text(slide, "Strength", 10.35, 7.8, 1.1, 0.24, size=10, color="gold2", bold=True, margin=0)
    add_text(slide, "works with one rated book and gives a concrete reason", 11.65, 7.8, 6.5, 0.24, size=11, color="paper", margin=0)
    add_text(slide, "Trade-off", 10.35, 8.35, 1.1, 0.24, size=10, color="gold2", bold=True, margin=0)
    add_text(slide, "can over-focus on known themes without reader-pattern discovery", 11.65, 8.35, 6.5, 0.4, size=11, color="paper", margin=0)


def slide_08(prs):
    slide = add_slide(prs, number=8, eyebrow="Collaborative Model", title="Readers reveal patterns that book metadata cannot", subtitle="Truncated SVD learns latent reader and book factors from the sparse rating matrix.")
    add_box(slide, 0.9, 2.7, 8.6, 6.7, fill="paper", line="line")
    add_card_header(slide, "Latent-factor view", 1.3, 3.1, 2.4, color="blue")
    # Matrix sketch
    add_text(slide, "readers", 1.3, 3.78, 1.0, 0.22, size=10, color="muted", bold=True, margin=0)
    add_text(slide, "books", 4.7, 3.45, 1.0, 0.22, size=10, color="muted", bold=True, margin=0)
    matrix_x, matrix_y = 2.0, 4.0
    for r in range(5):
        for c in range(7):
            val = (r * 3 + c * 2) % 5
            color = ["F2ECE1", "DCCFB8", "C5A059", "8E6D2C", "3D342A"][val]
            cell = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(matrix_x + c * 0.43), Inches(matrix_y + r * 0.43), Inches(0.35), Inches(0.35))
            set_fill(cell, color)
            set_line(cell, color, 0.4)
    add_text(slide, "mostly empty", 2.0, 6.35, 2.5, 0.25, size=10, color="muted", italic=True, margin=0)
    add_text(slide, "↓ SVD", 5.2, 5.0, 1.0, 0.32, size=14, color="blue", bold=True, align=PP_ALIGN.CENTER, margin=0)
    add_box(slide, 6.55, 4.15, 2.05, 1.75, fill="lavender", line="blue")
    add_text(slide, "40", 6.88, 4.42, 0.75, 0.55, size=30, color="blue", bold=True, font="Aptos Display", margin=0)
    add_text(slide, "latent factors", 7.52, 4.62, 0.85, 0.45, size=11, color="ink", bold=True, margin=0)
    add_text(slide, "taste dimensions learned from reader behavior", 6.88, 5.18, 1.35, 0.45, size=9, color="muted", margin=0)
    add_text(slide, "predicted rating", 2.0, 7.35, 1.55, 0.25, size=10, color="muted", bold=True, margin=0)
    add_text(slide, "μ + user bias + book bias + latent interaction", 3.65, 7.35, 4.7, 0.25, size=12, color="ink", font="Courier New", margin=0)
    add_text(slide, "Serendipity comes from discovering that readers who liked one book also liked another—even when the metadata does not overlap.", 1.3, 8.3, 7.5, 0.58, size=12, color="muted", italic=True, margin=0)

    add_box(slide, 9.95, 2.7, 8.85, 6.7, fill="ink", line="gold")
    add_card_header(slide, "Strengths and limits", 10.4, 3.1, 2.5, color="gold2")
    add_text(slide, "Strength", 10.4, 3.75, 1.2, 0.25, size=11, color="gold2", bold=True, margin=0)
    add_text(slide, "Finds cross-genre reader patterns and can predict ratings.", 11.85, 3.75, 6.0, 0.35, size=12, color="paper", margin=0)
    add_text(slide, "Risk", 10.4, 4.55, 1.2, 0.25, size=11, color="gold2", bold=True, margin=0)
    add_text(slide, "The matrix is 98.88% sparse; cold-start users have little evidence.", 11.85, 4.55, 6.0, 0.45, size=12, color="paper", margin=0)
    add_text(slide, "Fold-in", 10.4, 5.55, 1.2, 0.25, size=11, color="gold2", bold=True, margin=0)
    add_text(slide, "A new reader can be estimated from the ratings they have already submitted.", 11.85, 5.55, 6.0, 0.55, size=12, color="paper", margin=0)
    add_line(slide, 10.4, 6.7, 18.25, 6.7, color="ink2", width=1)
    add_text(slide, "Offline rating prediction", 10.4, 7.08, 2.8, 0.25, size=11, color="gold2", bold=True, margin=0)
    add_text(slide, "RMSE 0.9084   •   MAE 0.7174", 13.2, 7.08, 4.5, 0.25, size=15, color="paper", bold=True, font="Aptos Display", margin=0)
    add_text(slide, "measured on a reproducible 50,000-rating sample", 10.4, 7.62, 5.8, 0.25, size=10, color="ivory", margin=0)


def slide_09(prs):
    slide = add_slide(prs, number=9, eyebrow="Hybrid Ranking Logic", title="The production ranker adapts to how much we know", subtitle="BookWise gates the model by rating history, then balances reader fit with novelty and diversity.")
    add_box(slide, 0.85, 2.65, 18.25, 1.05, fill="ink", line="ink")
    add_text(slide, "base_hybrid", 1.25, 2.98, 1.5, 0.27, size=12, color="gold2", bold=True, font="Courier New", margin=0)
    add_text(slide, "= α × collaborative + (1 − α) × content", 3.05, 2.98, 5.4, 0.27, size=14, color="paper", bold=True, font="Courier New", margin=0)
    add_text(slide, "final score", 10.1, 2.98, 1.2, 0.27, size=12, color="gold2", bold=True, font="Courier New", margin=0)
    add_text(slide, "= 0.85 × base_hybrid + 0.15 × novelty", 11.7, 2.98, 6.4, 0.27, size=14, color="paper", bold=True, font="Courier New", margin=0)

    gates = [
        ("0 ratings", "Popularity + genre hooks", "Useful immediately", "gold"),
        ("1–2 ratings", "Content affinity", "A single signal can still personalize", "burgundy"),
        ("3+ ratings", "Hybrid ensemble", "Collaborative patterns join the ranking", "green"),
    ]
    for i, (history, mode, note, accent) in enumerate(gates):
        x = 1.0 + i * 6.0
        add_box(slide, x, 4.25, 5.45, 3.2, fill="paper", line="line")
        add_text(slide, history, x + 0.35, 4.62, 2.0, 0.34, size=18, color=accent, bold=True, font="Aptos Display", margin=0)
        add_text(slide, mode, x + 0.35, 5.25, 4.5, 0.32, size=14, color="ink", bold=True, margin=0)
        add_text(slide, note, x + 0.35, 5.85, 4.55, 0.5, size=11, color="muted", margin=0)
        add_pill(slide, "cold-start gate" if i < 2 else "full ranker", x + 0.35, 6.75, 1.45, fill="sand", color=accent, size=8)
        if i < 2:
            add_line(slide, x + 5.45, 5.85, x + 5.92, 5.85, color="gold", width=1.5, arrow=True)

    add_box(slide, 1.0, 8.1, 17.5, 1.3, fill="sand", line="line")
    add_text(slide, "Selection layer", 1.35, 8.43, 1.5, 0.25, size=11, color="gold", bold=True, margin=0)
    add_text(slide, "Greedy MMR-style selection penalizes repeated genres and rewards metadata diversity so a shelf does not collapse into ten near-duplicates.", 3.1, 8.39, 14.6, 0.45, size=12, color="ink", margin=0)


def slide_10(prs):
    slide = add_slide(prs, number=10, eyebrow="Evaluation Results", title="The hybrid model leads the offline ranking benchmark", subtitle="Held-out evaluation: 250 eligible users and 50,000 sampled ratings. Metrics are not a claim over the full test set.")
    add_box(slide, 0.85, 2.62, 10.2, 6.98, fill="paper", line="line")
    add_card_header(slide, "Ranking quality", 1.25, 3.0, 2.3, color="green")
    add_text(slide, "Higher is better", 8.3, 3.02, 2.1, 0.22, size=9, color="muted", align=PP_ALIGN.RIGHT, margin=0)
    add_text(slide, "Precision@5", 1.35, 3.65, 2.1, 0.25, size=11, color="muted", bold=True, margin=0)
    add_text(slide, "NDCG@10", 1.35, 6.25, 2.1, 0.25, size=11, color="muted", bold=True, margin=0)
    models = [
        ("Popularity", 0.0912, 0.2338, "gold"),
        ("Content-based", 0.1544, 0.3437, "burgundy"),
        ("Collaborative", 0.0392, 0.1402, "blue"),
        ("Hybrid", 0.1936, 0.4395, "green"),
    ]
    for i, (label, p5, ndcg, color) in enumerate(models):
        y = 4.05 + i * 0.56
        add_bar(slide, label, p5, 0.22, 1.35, y, 8.5, color=color)
        y2 = 6.65 + i * 0.56
        add_bar(slide, label, ndcg, 0.5, 1.35, y2, 8.5, color=color)
    add_line(slide, 1.35, 5.95, 10.5, 5.95, color="line", width=1)
    add_text(slide, "Hybrid also reaches 11.33% of the catalogue in the sampled evaluation.", 1.35, 8.85, 8.7, 0.35, size=12, color="green", bold=True, margin=0)

    add_box(slide, 11.45, 2.62, 7.4, 6.98, fill="ink", line="gold")
    add_card_header(slide, "Read the result correctly", 11.9, 3.0, 3.2, color="gold2")
    add_text(slide, "0.1936", 11.9, 3.75, 2.0, 0.6, size=32, color="gold2", bold=True, font="Aptos Display", margin=0)
    add_text(slide, "Hybrid Precision@5", 14.0, 3.98, 3.6, 0.25, size=13, color="paper", bold=True, margin=0)
    add_bullets(slide, [
        ("Best ranking quality. ", "Hybrid beats popularity, content-only and SVD on Precision@5 and NDCG@10."),
        ("Not a probability. ", "Scores rank candidates; they do not mean a reader has a 19.36% chance of liking a book."),
        ("Sampled benchmark. ", "The report intentionally uses a reproducible slice for ranking metrics."),
        ("Coverage trade-off. ", "Diversity is measured alongside relevance; the hybrid recommends 1,133 unique titles in the sample."),
    ], 11.9, 5.05, 6.25, 3.35, size=12, color="ivory", bullet_color="gold2", gap=0.12)


def slide_11(prs):
    slide = add_slide(prs, number=11, eyebrow="Personalization and Cold Start", title="Personalization starts before the reader has a profile", subtitle="The discovery feed changes its strategy as the evidence grows.")
    add_box(slide, 0.9, 2.72, 18.2, 1.0, fill="sand", line="line")
    add_text(slide, "new reader", 1.3, 3.05, 1.2, 0.25, size=11, color="muted", bold=True, margin=0)
    add_line(slide, 2.7, 3.17, 17.15, 3.17, color="gold", width=2)
    for x, label, color in [(3.25, "0", "gold"), (9.3, "1–2", "burgundy"), (15.25, "3+", "green")]:
        circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(2.9), Inches(0.56), Inches(0.56))
        set_fill(circ, color)
        set_line(circ, color)
        add_text(slide, label, x, 3.04, 0.56, 0.2, size=10, color="paper", bold=True, align=PP_ALIGN.CENTER, margin=0)
    add_text(slide, "ratings accumulated", 17.25, 3.05, 1.25, 0.25, size=10, color="muted", bold=True, align=PP_ALIGN.RIGHT, margin=0)

    stages = [
        ("0 ratings", "Cold start", "Popularity baseline", ["popular shelf", "genre hooks", "no friction"], "gold"),
        ("1–2 ratings", "Warm start", "Content affinity", ["similar themes", "similar authors", "immediate signal"], "burgundy"),
        ("3+ ratings", "Established", "Hybrid ensemble", ["CF + content", "novelty / diversity", "reader feedback"], "green"),
    ]
    for i, (count, label, engine, points, accent) in enumerate(stages):
        x = 0.9 + i * 6.1
        add_box(slide, x, 4.35, 5.45, 4.65, fill="paper", line="line")
        add_text(slide, count, x + 0.35, 4.75, 1.5, 0.3, size=17, color=accent, bold=True, font="Aptos Display", margin=0)
        add_text(slide, label, x + 2.1, 4.8, 2.3, 0.25, size=11, color="muted", align=PP_ALIGN.RIGHT, margin=0)
        add_line(slide, x + 0.35, 5.35, x + 5.1, 5.35, color="line", width=1)
        add_text(slide, engine, x + 0.35, 5.65, 4.2, 0.32, size=15, color="ink", bold=True, margin=0)
        for j, point in enumerate(points):
            add_pill(slide, point, x + 0.35, 6.35 + j * 0.55, 2.0 if len(point) < 15 else 2.8, fill="sand", color=accent, size=9, line="line")
    add_text(slide, "The reader can also shape the shelf directly: favorite genres, disliked genres, novelty, diversity, feedback events and alpha (similar-reader influence).", 1.1, 9.45, 17.5, 0.42, size=13, color="ink", italic=True, align=PP_ALIGN.CENTER, margin=0)


def slide_12(prs):
    slide = add_slide(prs, number=12, eyebrow="ML Insights", title="The app exposes how the recommender behaves", subtitle="BookWise includes an ML Insights page for model comparison, genre prediction and reader taste clustering.")
    add_box(slide, 0.85, 2.65, 8.65, 6.95, fill="ink", line="gold")
    add_card_header(slide, "Genre prediction", 1.3, 3.05, 2.4, color="gold2")
    add_text(slide, "11 canonical genres", 1.3, 3.62, 3.4, 0.38, size=22, color="paper", bold=True, font="Aptos Display", margin=0)
    add_text(slide, "Features: title + authors + tags + optional description", 1.3, 4.15, 6.8, 0.3, size=11, color="ivory", margin=0)
    add_text(slide, "Best benchmarked model", 1.3, 4.95, 2.35, 0.25, size=10, color="gold2", bold=True, margin=0)
    add_text(slide, "Random Forest", 3.8, 4.9, 2.5, 0.35, size=17, color="paper", bold=True, font="Aptos Display", margin=0)
    add_text(slide, "97.96% accuracy", 1.3, 5.55, 2.55, 0.45, size=25, color="gold2", bold=True, font="Aptos Display", margin=0)
    add_text(slide, "0.8732 macro-F1", 4.2, 5.65, 2.6, 0.3, size=14, color="paper", bold=True, margin=0)
    add_line(slide, 1.3, 6.45, 8.9, 6.45, color="ink2", width=1)
    add_bullets(slide, [
        ("Leakage-safe. ", "primary_genre is the target, never a prediction feature."),
        ("Conservative taxonomy. ", "unknown tags are reported, not silently promoted to new genres."),
        ("Caveat. ", "rare classes such as Adventure and Horror have unstable per-class scores."),
    ], 1.3, 6.85, 7.35, 1.95, size=11, color="ivory", bullet_color="gold2", gap=0.08)

    add_box(slide, 9.95, 2.65, 8.85, 6.95, fill="paper", line="line")
    add_card_header(slide, "Taste clusters", 10.4, 3.05, 2.0, color="green")
    add_text(slide, "K-Means groups app readers by the genres they rate highly.", 10.4, 3.65, 7.1, 0.35, size=13, color="ink", margin=0)
    cluster_labels = [("Fantasy-led", 0.78, "burgundy"), ("Classics & fiction", 0.62, "gold"), ("Mystery & thriller", 0.48, "blue"), ("Nonfiction", 0.35, "green")]
    for i, (label, value, color) in enumerate(cluster_labels):
        y = 4.45 + i * 0.78
        add_text(slide, label, 10.4, y, 2.5, 0.24, size=11, color="ink", margin=0)
        add_box(slide, 13.0, y + 0.04, 4.2, 0.18, fill="line", line="line")
        add_box(slide, 13.0, y + 0.04, 4.2 * value, 0.18, fill=color, line=color)
        add_text(slide, f"{int(value * 100)}%", 17.45, y - 0.02, 0.6, 0.25, size=10, color=color, bold=True, margin=0)
    add_line(slide, 10.4, 7.7, 18.2, 7.7, color="line", width=1)
    add_text(slide, "The cluster becomes a soft signal—not a hard label—in the personalized feed.", 10.4, 8.05, 7.3, 0.45, size=13, color="green", bold=True, margin=0)


def slide_13(prs):
    slide = add_slide(prs, number=13, eyebrow="Implementation and Stack", title="A practical full-stack build, not just a notebook", subtitle="The project connects offline model work to a usable reader-facing application.")
    columns = [
        ("Frontend", ["React 19", "Vite 8", "Tailwind CSS 4", "Lucide icons"], "gold"),
        ("Backend", ["Python 3.12", "Flask 3", "SQLAlchemy 2", "Gunicorn"], "burgundy"),
        ("Data / ML", ["Goodbooks-10k", "TF-IDF", "Truncated SVD", "K-Means + genre models"], "blue"),
        ("Quality / deploy", ["unittest", "Vitest + Testing Library", "Oxlint", "Render + PostgreSQL"], "green"),
    ]
    for i, (label, items, accent) in enumerate(columns):
        x = 0.85 + i * 4.65
        add_box(slide, x, 2.75, 4.15, 3.6, fill="paper", line="line")
        add_text(slide, label, x + 0.3, 3.1, 3.3, 0.32, size=17, color=accent, bold=True, font="Aptos Display", margin=0)
        add_line(slide, x + 0.3, 3.62, x + 3.8, 3.62, color="line", width=1)
        for j, item in enumerate(items):
            add_text(slide, "•", x + 0.32, 4.0 + j * 0.48, 0.2, 0.2, size=12, color=accent, bold=True, margin=0)
            add_text(slide, item, x + 0.62, 4.0 + j * 0.48, 3.0, 0.22, size=12, color="ink", margin=0)

    add_box(slide, 0.85, 6.85, 18.05, 2.55, fill="ink", line="gold")
    add_card_header(slide, "Reader-owned capabilities", 1.3, 7.22, 3.6, color="gold2")
    capabilities = [
        "secure cookie-based auth",
        "ratings, favorites and shelves",
        "reading progress and notes",
        "reviews and comments",
        "private lists + public share links",
        "feedback and experiments",
    ]
    for i, cap in enumerate(capabilities):
        x = 1.3 + (i % 3) * 5.45
        y = 7.75 + (i // 3) * 0.62
        add_pill(slide, cap, x, y, 3.6 if len(cap) > 25 else 3.1, fill="ink2", color="gold2", size=9, line="gold")
    add_text(slide, "The recommendation is only one part of the product; the reader can keep using the app after the ranked shelf is generated.", 1.3, 9.0, 16.8, 0.26, size=11, color="ivory", italic=True, align=PP_ALIGN.CENTER, margin=0)


def slide_14(prs):
    slide = add_slide(prs, number=14, eyebrow="Evaluation and Limitations", title="What we measured—and what we still need to measure", subtitle="The current results are useful evidence, not a claim that the system is finished.")
    add_box(slide, 0.85, 2.65, 8.6, 6.95, fill="paper", line="line")
    add_card_header(slide, "Evidence we have", 1.3, 3.05, 2.8, color="green")
    add_bullets(slide, [
        ("Offline ranking benchmark. ", "Hybrid leads the tracked Precision@5, Recall@10 and NDCG@10 metrics."),
        ("Dataset audit. ", "10,000 books, 53,424 users, 5.98M ratings, 98.88% sparsity."),
        ("Genre benchmark. ", "Random Forest reaches 97.96% accuracy on the held-out genre split."),
        ("Product checks. ", "Frontend, backend and browser smoke tests cover major user flows."),
    ], 1.3, 3.7, 7.3, 3.2, size=13, color="ink", bullet_color="green", gap=0.14)
    add_pill(slide, "measured", 1.3, 8.35, 1.05, fill="sand", color="green", size=9)
    add_text(slide, "Evidence is tied to the artifacts and reproducible commands in the repository.", 2.55, 8.4, 6.0, 0.25, size=11, color="muted", margin=0)

    add_box(slide, 9.95, 2.65, 8.85, 6.95, fill="ink", line="gold")
    add_card_header(slide, "Open limitations", 10.4, 3.05, 2.7, color="gold2")
    add_bullets(slide, [
        ("Sampled evaluation. ", "Ranking metrics use 250 eligible users and 50,000 held-out ratings."),
        ("No live catalogue yet. ", "The data is a static academic dataset, not a production-scale feed."),
        ("Rare genre classes. ", "Some per-class genre scores are unstable because support is tiny."),
        ("Weak multi-label targets. ", "Secondary genre labels are metadata-derived until reviewed annotations exist."),
        ("Offline ≠ online. ", "Human judgments, long-term retention and fairness still need study."),
    ], 10.4, 3.7, 7.5, 3.8, size=12, color="ivory", bullet_color="gold2", gap=0.1)
    add_pill(slide, "do not overclaim", 10.4, 8.4, 1.55, fill="ink2", color="gold2", size=9, line="gold")
    add_text(slide, "A good benchmark is a starting point for better experiments.", 12.2, 8.46, 5.7, 0.25, size=11, color="ivory", italic=True, margin=0)


def slide_15(prs):
    slide = add_slide(prs, bg="ink")
    add_text(slide, "15", 0.8, 0.72, 0.5, 0.3, size=13, color="gold2", bold=True, margin=0)
    add_text(slide, "CONCLUSION AND FUTURE SCOPE", 1.42, 0.72, 5.4, 0.26, size=10, color="muted", bold=True, margin=0)
    add_text(slide, "BookWise turns a book catalogue into a reading companion.", 0.8, 1.58, 13.8, 1.0, size=38, color="paper", bold=True, font="Aptos Display", margin=0)
    add_text(slide, "It starts with popularity, learns from metadata and reader behavior, and keeps the reader in control of the result.", 0.85, 2.88, 12.6, 0.55, size=16, color="ivory", margin=0)

    add_box(slide, 0.85, 4.2, 8.4, 4.45, fill="ink2", line="gold")
    add_card_header(slide, "What is working now", 1.3, 4.58, 3.2, color="gold2")
    add_bullets(slide, [
        "A 10,000-book discovery surface",
        "Four complementary recommendation strategies",
        "Cold-start fallback and explainable ranking",
        "Reader library, feedback and shareable lists",
        "Offline evaluation and ML inspection tools",
    ], 1.3, 5.18, 7.2, 2.75, size=14, color="paper", bullet_color="gold2", gap=0.12)

    add_box(slide, 9.75, 4.2, 9.1, 4.45, fill="paper", line="gold")
    add_card_header(slide, "Next experiments", 10.2, 4.58, 3.0, color="burgundy")
    add_bullets(slide, [
        "Evaluate on the full held-out test set",
        "Add human relevance judgments and A/B tests",
        "Expand taxonomy and reviewed multi-label data",
        "Measure fairness, diversity and long-term retention",
        "Connect to larger or live book catalogues",
    ], 10.2, 5.18, 7.8, 2.75, size=14, color="ink", bullet_color="burgundy", gap=0.12)

    add_text(slide, "Thank you", 0.85, 9.2, 3.0, 0.45, size=24, color="gold2", bold=True, font="Aptos Display", margin=0)
    add_text(slide, "Questions?", 15.75, 9.28, 3.0, 0.35, size=20, color="paper", bold=True, align=PP_ALIGN.RIGHT, margin=0)
    add_footer(slide, "BookWise • Intelligent Book Discovery")


def build_presentation():
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    prs.core_properties.title = "BookWise - Intelligent Book Discovery & Recommendation"
    prs.core_properties.subject = "Project presentation for the BookWise recommendation system"
    prs.core_properties.author = "BookWise Project Team"
    prs.core_properties.keywords = "recommendation systems, books, hybrid recommender, machine learning"

    slide_01(prs)
    slide_02(prs)
    slide_03(prs)
    slide_04(prs)
    slide_05(prs)
    slide_06(prs)
    slide_07(prs)
    slide_08(prs)
    slide_09(prs)
    slide_10(prs)
    slide_11(prs)
    slide_12(prs)
    slide_13(prs)
    slide_14(prs)
    slide_15(prs)

    prs.save(OUTPUT)
    print(f"Saved {OUTPUT}")


if __name__ == "__main__":
    build_presentation()
